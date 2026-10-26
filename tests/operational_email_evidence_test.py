import datetime as dt
import hashlib
import importlib.util
from pathlib import Path
import tempfile
import unittest

import openpyxl

spec = importlib.util.spec_from_file_location('evidence', Path(__file__).parents[1] /
                                             'scripts/operational-email-evidence.py')
evidence = importlib.util.module_from_spec(spec)
spec.loader.exec_module(evidence)


class EvidenceTests(unittest.TestCase):
    def profile(self, date, title='Reporte Diario', declared=None):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.xlsx'
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = title
            sheet['A2'] = 'Fecha del reporte'
            sheet['E2'] = date
            future = book.create_sheet('Plan futuro')
            future.sheet_state = 'hidden'
            future['E2'] = dt.datetime(2030, 1, 1)
            future['A1'] = '=1/0'
            book.save(path)
            return evidence.workbook_profile(path, title, declared)

    def test_internal_date_never_comes_from_future_plan_or_receipt(self):
        result = self.profile(dt.datetime(2026, 10, 20))
        self.assertEqual(result['operational_date'], '2026-10-20')
        self.assertEqual(result['sheets'][1]['state'], 'hidden')
        self.assertEqual(result['sheets'][1]['formula_count'], 1)
        self.assertFalse(result['canonical_promotion_allowed'])

    def test_raiz_sheet_name_cannot_override_internal_date(self):
        result = self.profile(dt.datetime(2026, 10, 22), '26-10')
        self.assertEqual(result['operational_date'], '2026-10-22')
        self.assertIn('sheet_date_mismatch', result['issues'])

    def test_explicit_source_date_conflict(self):
        result = self.profile(dt.datetime(2026, 10, 20), declared='2026-10-26')
        self.assertIn('declared_internal_date_mismatch', result['issues'])

    def test_invalid_or_uncached_formula_date_stays_unresolved(self):
        for date in ('31/02/2026', '=TODAY()', None):
            with self.subTest(date=date):
                self.assertIsNone(self.profile(date)['operational_date'])

    def test_receipt_is_separate_and_late_report_is_not_automatically_invalid(self):
        blob = b'original'
        manifest = {'organization_id': evidence.ORG, 'email_sender': 'Taller <taller@lapatagua.cl>',
                    'source_sha256': hashlib.sha256(blob).hexdigest(), 'source_size_bytes': len(blob),
                    'email_sent_at': '2026-10-27T01:00:00+00:00', 'source_profile': {}}
        profile = evidence.validate_manifest(manifest, blob, self.profile(dt.datetime(2026, 10, 20)))
        self.assertEqual(profile['email_sent_local_date'], '2026-10-26')
        self.assertEqual(profile['receipt_lag_days'], 6)
        self.assertEqual(profile['issues'], [])
        with self.assertRaisesRegex(ValueError, 'source_hash_mismatch'):
            evidence.validate_manifest(manifest, b'changed', profile)

    def test_same_filename_preserves_changed_versions_and_repeated_transmissions(self):
        manifest = {'id': 'new', 'source_sha256': 'a'}
        links = evidence.version_links(manifest, [
            {'id': 'old1', 'gmail_message_id': 'm1', 'source_sha256': 'b'},
            {'id': 'old2', 'gmail_message_id': 'm2', 'source_sha256': 'a'}])
        self.assertEqual([x['relation'] for x in links], ['different_version', 'same_bytes'])
        self.assertEqual(len(links), 2)

    def test_statistics_cutoff_uses_printed_header(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'statistics.xlsx'
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = 'OCTUBRE 2026'
            sheet['B2'] = 'ESTADISTICA DE EXTRACCION A PLANTA OCTUBRE 2026'
            sheet['H3'] = 'FECHA:'
            sheet['I3'] = dt.datetime(2026, 10, 25)
            book.save(path)
            result = evidence.workbook_profile(path, sheet.title, '2026-10-25',
                                                'production_statistics')
            self.assertEqual(result['operational_date'], '2026-10-25')
            self.assertEqual(result['date_cell'], 'I3')
            self.assertEqual(result['issues'], [])

    def test_new_conflict_reaches_existing_review_flags_without_erasing_authority(self):
        old = {'source_profile': {'authority_correction_message_id': 'authority',
                                  'workflow_status': 'needs_report_correction',
                                  'promotion': 'domain_review_on_conflict',
                                  'promoted_to_canonical': False}}
        profile = {'issues': ['sheet_date_mismatch']}
        merged = evidence.merged_profile(old, profile, 'private/object.xlsx')
        self.assertTrue(merged['requires_review'])
        self.assertTrue(merged['date_mismatch'])
        self.assertEqual(merged['promotion'], 'human_review_required')
        self.assertEqual(merged['authority_correction_message_id'], 'authority')
        self.assertEqual(merged['workflow_status'], 'needs_report_correction')
        self.assertFalse(merged['promoted_to_canonical'])


if __name__ == '__main__':
    unittest.main()
