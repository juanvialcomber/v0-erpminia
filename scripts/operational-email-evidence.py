#!/usr/bin/env python3
"""Bounded recovery of existing email manifests. Never writes canonical facts.

Requires openpyxl. Default is a read-only plan. Input is a private JSON array of
{manifest_id, local_path, operational_sheet?, declared_coverage_date?, report_kind?} records.
Credentials come from SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY only.
"""
import argparse
import datetime as dt
import hashlib
import json
import os
from pathlib import Path
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
import uuid
from zoneinfo import ZoneInfo

PROJECT_HOST = 'ttlptyheuqeotadtcbaw.supabase.co'
ORG = '2bd7fe06-8e4f-4a3a-b261-e3f5d8aa3dee'
BUCKET = 'module-documents'
MAX_BYTES = 20 * 1024 * 1024


def iso_date(value):
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, str):
        for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y'):
            try:
                return dt.datetime.strptime(value.strip(), fmt).date().isoformat()
            except ValueError:
                pass
    return None


def workbook_profile(path, operational_sheet=None, declared_coverage_date=None, report_kind='equipment'):
    import openpyxl
    cached = openpyxl.load_workbook(path, data_only=True, read_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
    sheets = []
    try:
        for ws in formulas:
            count = sum(c.data_type == 'f' for row in ws for c in row)
            sheets.append({'name': ws.title, 'state': ws.sheet_state,
                           'rows': ws.max_row, 'columns': ws.max_column,
                           'formula_count': count})
        result = {'sheets': sheets, 'declared_coverage_date': None,
                  'operational_date': None, 'date_basis': 'unresolved',
                  'issues': [], 'canonical_promotion_allowed': False}
        if declared_coverage_date:
            if not iso_date(declared_coverage_date):
                raise ValueError('invalid_declared_coverage_date')
            result['declared_coverage_date'] = iso_date(declared_coverage_date)
        # A report header is authoritative evidence of the date written in the
        # file, not proof that the date is correct. Never take max(workbook dates).
        if operational_sheet:
            if operational_sheet not in cached.sheetnames:
                raise ValueError('operational_sheet_missing')
            ws = cached[operational_sheet]
            label_cell, date_cell, label = ('A2', 'E2', 'fecha del reporte')
            if report_kind == 'production_statistics':
                if 'ESTADISTICA' not in str(ws['B2'].value or '').upper():
                    raise ValueError('unrecognized_statistics_header')
                label_cell, date_cell, label = ('H3', 'I3', 'fecha:')
            elif report_kind != 'equipment':
                raise ValueError('unsupported_report_kind')
            if str(ws[label_cell].value or '').strip().casefold() != label:
                raise ValueError('unrecognized_report_date_header')
            raw = ws[date_cell].value
            result.update(operational_sheet=operational_sheet,
                          date_cell=date_cell, date_raw=str(raw),
                          operational_date=iso_date(raw), date_basis='internal_header')
            if not result['operational_date']:
                result['issues'].append('missing_or_invalid_internal_date')
            else:
                internal = dt.date.fromisoformat(result['operational_date'])
                match = re.fullmatch(r'(\d{1,2})[-/](\d{1,2})(?:[-/](\d{4}))?', operational_sheet)
                if match:
                    try:
                        named = dt.date(int(match[3] or internal.year), int(match[2]), int(match[1]))
                        if named != internal:
                            result['issues'].append('sheet_date_mismatch')
                    except ValueError:
                        result['issues'].append('invalid_sheet_date')
                declared = result['declared_coverage_date']
                if declared and declared != internal.isoformat():
                    result['issues'].append('declared_internal_date_mismatch')
        else:
            # Production/stock books mix historical actuals and future plans.
            # A sender-declared cutoff is retained but not certified by this tool.
            result['date_basis'] = 'declared_coverage' if declared_coverage_date else 'unresolved'
            result['issues'].append('operational_coverage_requires_domain_review')
        return result
    finally:
        cached.close()
        formulas.close()


def validate_manifest(manifest, blob, profile):
    if manifest['organization_id'] != ORG:
        raise ValueError('organization_mismatch')
    sender = manifest['email_sender'].strip().lower()
    if not re.search(r'@lapatagua\.cl(?:>|\s|$)', sender):
        raise ValueError('unverified_sender_domain')
    if hashlib.sha256(blob).hexdigest() != manifest['source_sha256']:
        raise ValueError('source_hash_mismatch')
    if len(blob) != manifest['source_size_bytes']:
        raise ValueError('source_size_mismatch')
    sent = dt.datetime.fromisoformat(manifest['email_sent_at']).astimezone(ZoneInfo('America/Santiago')).date()
    internal = profile['operational_date']
    if internal and dt.date.fromisoformat(internal) > sent:
        profile['issues'].append('future_report_date')
    # A lag alone is not a date error: late reports are valid evidence. Only
    # explicit assertions/conflicts can establish a discrepancy.
    profile['email_sent_local_date'] = sent.isoformat()
    profile['receipt_lag_days'] = (sent - dt.date.fromisoformat(internal)).days if internal else None
    old = manifest.get('source_profile') or {}
    if old.get('date_mismatch') or old.get('date_discrepancy'):
        profile['issues'].append('existing_date_discrepancy')
    if old.get('authority_correction_message_id'):
        profile['authority_message_id'] = old['authority_correction_message_id']
        profile['issues'].append('authority_correction_pending')
    profile['issues'] = sorted(set(profile['issues']))
    return profile


def version_links(manifest, previous):
    """Each transmission keeps its identity, even when bytes repeat."""
    return [{'manifest_id': row['id'], 'gmail_message_id': row['gmail_message_id'],
             'relation': 'same_bytes' if row['source_sha256'] == manifest['source_sha256']
             else 'different_version', 'source_sha256': row['source_sha256']}
            for row in previous if row['id'] != manifest['id']]


class Client:
    def __init__(self):
        self.url = os.environ['SUPABASE_URL'].rstrip('/')
        parsed = urllib.parse.urlparse(self.url)
        if parsed.scheme != 'https' or parsed.netloc != PROJECT_HOST or parsed.path:
            raise ValueError('wrong_project_endpoint')
        key = os.environ['SUPABASE_SERVICE_ROLE_KEY']
        self.headers = {'apikey': key, 'Authorization': 'Bearer ' + key}

    def request(self, path, method='GET', body=None, content_type='application/json'):
        headers = {**self.headers, 'Content-Type': content_type,
                   'Prefer': 'return=representation'}
        req = urllib.request.Request(self.url + path, data=body, headers=headers, method=method)
        try:
            with urllib.request.urlopen(req, timeout=30) as response:
                return response.read(MAX_BYTES + 1)
        except urllib.error.HTTPError as error:
            # Never include request URLs, provider body or credentials in errors.
            raise RuntimeError('request_http_' + str(error.code)) from None

    def manifest(self, manifest_id):
        uuid.UUID(manifest_id)
        query = urllib.parse.urlencode({'id': 'eq.' + manifest_id,
                                       'organization_id': 'eq.' + ORG, 'limit': '1'})
        rows = json.loads(self.request('/rest/v1/motil_email_source_manifest?' + query))
        if len(rows) != 1:
            raise ValueError('manifest_not_found')
        return rows[0]

    def preserve(self, manifest, blob):
        bucket = json.loads(self.request('/storage/v1/bucket/' + BUCKET))
        if bucket.get('public') is not False:
            raise ValueError('private_bucket_required')
        path = f"operational-email/{ORG}/{manifest['id']}/{manifest['source_sha256']}.xlsx"
        encoded = urllib.parse.quote(path, safe='/')
        get_path = '/storage/v1/object/authenticated/' + BUCKET + '/' + encoded
        try:
            stored = self.request(get_path)
        except RuntimeError as error:
            if str(error) not in ('request_http_400', 'request_http_404'):
                raise
            # No upsert: existing originals must never be overwritten. A race
            # causing 409 is safe; verify the winning object below.
            try:
                self.request('/storage/v1/object/' + BUCKET + '/' + encoded, 'POST', blob,
                             'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            except RuntimeError as upload_error:
                if str(upload_error) != 'request_http_409':
                    raise
            stored = self.request(get_path)
        if stored != blob:
            raise ValueError('stored_original_verification_failed')
        return path

    def previous_versions(self, manifest):
        params = {'organization_id': 'eq.' + ORG, 'filename': 'eq.' + manifest['filename'],
                  'email_sent_at': 'lte.' + manifest['email_sent_at'],
                  'select': 'id,gmail_message_id,source_sha256',
                  'order': 'email_sent_at.desc', 'limit': '6'}
        rows = json.loads(self.request('/rest/v1/motil_email_source_manifest?' +
                                      urllib.parse.urlencode(params)))
        return version_links(manifest, rows)

    def patch(self, manifest, profile, storage_path):
        current = dict(manifest.get('source_profile') or {})
        current.update(raw_original_in_motil_storage=True, sha256_verified=True,
                       storage_bucket=BUCKET, storage_path=storage_path,
                       operational_validation_v1=profile)
        if current.get('intake_exception') == 'raw_original_not_in_motil_storage':
            current.pop('intake_exception')
            current['intake_exception_open'] = False
        # Preserve all unrelated annotations, source dates, authority and states.
        params = {'id': 'eq.' + manifest['id'], 'organization_id': 'eq.' + ORG,
                  'updated_at': 'eq.' + manifest['updated_at']}
        rows = json.loads(self.request('/rest/v1/motil_email_source_manifest?' +
                                      urllib.parse.urlencode(params), 'PATCH',
                                      json.dumps({'source_profile': current}).encode()))
        if len(rows) != 1:
            raise ValueError('manifest_changed_concurrently_replan')
        readback = self.manifest(manifest['id'])
        if readback['source_profile'] != current:
            raise ValueError('manifest_readback_mismatch')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('batch')
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    items = json.loads(Path(args.batch).read_text())
    if not isinstance(items, list) or not 1 <= len(items) <= 10:
        raise ValueError('batch_must_contain_1_to_10_records')
    if len({x['manifest_id'] for x in items}) != len(items):
        raise ValueError('duplicate_manifest_in_batch')
    client = Client()
    for item in items:
        manifest = client.manifest(item['manifest_id'])
        path = Path(item['local_path'])
        if path.stat().st_size > MAX_BYTES:
            raise ValueError('source_too_large')
        blob = path.read_bytes()
        profile = validate_manifest(manifest, blob, workbook_profile(
            path, item.get('operational_sheet'), item.get('declared_coverage_date'),
            item.get('report_kind', 'equipment')))
        profile['previous_transmissions'] = client.previous_versions(manifest)
        if args.apply:
            storage_path = client.preserve(manifest, blob)
            client.patch(manifest, profile, storage_path)
        print(json.dumps({'manifest_id': manifest['id'], 'applied': args.apply,
                          'operational_date': profile['operational_date'],
                          'date_basis': profile['date_basis'], 'issues': profile['issues'],
                          'sheets': len(profile['sheets'])}))


if __name__ == '__main__':
    try:
        main()
    except Exception as exc:
        # Avoid accidental disclosure of local filenames or connection secrets.
        code = str(exc) if isinstance(exc, (ValueError, RuntimeError)) else type(exc).__name__
        print(json.dumps({'error': code}), file=sys.stderr)
        sys.exit(1)
